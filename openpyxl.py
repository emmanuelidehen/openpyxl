#!/usr/bin/env 
# -python3*- coding: utf-8 -*-
"""
Created on Mon Aug 24 19:01:54 2020

@author: emmanuelidehen
"""
import pandas as pd
import numpy as np
import openpyxl
from openpyxl import load_workbook
from openpyxl import worksheet
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, PieChart, Reference



wb = load_workbook('Numerical_Analysis_Emmanuel_idehen_Report.xlsx')

ws = wb.active 
chart = BarChart()
data = Reference(ws, min_row=8, min_col = 1, max_col=13, max_row=13 )

chart.add_data(data, titles_from_data=True)
ws.add_chart(chart, 'B14')
wb.save('line.xlsx')
